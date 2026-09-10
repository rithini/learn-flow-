from typing import List, Dict, Any
from app.document_processing.extractors.base import BaseExtractor
from app.core.logging import logger


class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: str) -> List[Dict[str, Any]]:
        units: List[Dict[str, Any]] = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text)
                combined = "\n".join(slide_texts)
                if combined.strip():
                    units.append({
                        "content": combined,
                        "locator": {"slide_number": idx + 1, "total_slides": len(prs.slides)},
                    })
        except ImportError:
            logger.warning("python-pptx not available. Reading fallback.")
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                units.append({"content": f.read(), "locator": {"slide_number": 1}})
        except Exception as e:
            logger.error(f"Error extracting PPTX from {file_path}: {e}")
            raise e
        return units


class DOCXExtractor(BaseExtractor):
    def extract(self, file_path: str) -> List[Dict[str, Any]]:
        units: List[Dict[str, Any]] = []
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            if full_text.strip():
                units.append({
                    "content": full_text,
                    "locator": {"section": "main_body", "paragraph_count": len(paragraphs)},
                })
        except ImportError:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                units.append({"content": f.read(), "locator": {"section": "main_body"}})
        except Exception as e:
            logger.error(f"Error extracting DOCX from {file_path}: {e}")
            raise e
        return units


class TXTExtractor(BaseExtractor):
    def extract(self, file_path: str) -> List[Dict[str, Any]]:
        units: List[Dict[str, Any]] = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
            units.append({
                "content": content,
                "locator": {"type": "plain_text"},
            })
        return units
